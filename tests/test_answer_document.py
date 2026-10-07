"""Local export tests only; fixture generation is not live AI or OAuth evidence."""

import io
import json

import pytest
from docx import Document
from test_frontend import local_server as bounded_loopback_server
from test_workflow import prepared_fixture

from classroomautowork.answer_document import (
    answer_blocks,
    create_answer_document,
    valid_answer_document,
)
from classroomautowork.errors import WorkflowError
from classroomautowork.local import atomic_json, sha256_file
from classroomautowork.review import finalize


def verified_package(tmp_path):
    package, review_path, draft, review = prepared_fixture(tmp_path)
    evidence_id = review["claim_checks"][0]["evidence_ids"][0]
    draft.write_text(
        f"# Offline report\n\n食品安全 [E:{evidence_id}]\n\n"
        "【本机审核说明】Private audit note.\n\n"
        "本人評価：未入力\n\n## 本机证据记录（不进入提交正文）\n\nPrivate evidence index.",
        encoding="utf-8",
    )
    review["requirement_checks"][0]["status"] = "needs_user"
    atomic_json(review_path, review)
    finalize(package, review_path, draft)
    atomic_json(
        package / "codex-generation.json", {"status": "completed", "model": "offline-fixture"}
    )
    return package


def test_verified_report_is_readable_and_reuses_exact_files_without_ai(tmp_path):
    package = verified_package(tmp_path)
    original_hash = sha256_file(package / "draft.md")
    result = create_answer_document(package)
    assert result["status"] == "needs_user"
    assert result["pending_requirements"]
    content = "\n".join(p.text for p in Document(package / "answer.docx").paragraphs)
    assert "食品安全" in content and "本人評価：未入力" in content
    assert "[E:" not in content and "Private evidence index" not in content
    assert "Private audit note" not in content
    assert Document(package / "answer.docx").paragraphs[0].style.name == "Title"
    assert not Document(package / "answer.docx").styles["Title"].element.xpath(".//w:pBdr")
    assert sha256_file(package / "draft.md") == original_hash
    timestamp = (package / "answer.docx").stat().st_mtime_ns
    assert create_answer_document(package) == result
    assert (package / "answer.docx").stat().st_mtime_ns == timestamp
    (package / "answer.docx").write_bytes(b"broken")
    assert valid_answer_document(package) is None
    assert create_answer_document(package)["files"]["answer.docx"] == sha256_file(
        package / "answer.docx"
    )


def test_changed_or_unfinished_ai_output_is_never_exported_as_verified(tmp_path):
    package = verified_package(tmp_path)
    create_answer_document(package)
    atomic_json(package / "codex-generation.json", {"status": "running"})
    assert valid_answer_document(package) is None
    with pytest.raises(WorkflowError, match="尚未完成核验"):
        create_answer_document(package)
    atomic_json(package / "codex-generation.json", {"status": "completed"})
    (package / "draft.md").write_text("unverified modification", encoding="utf-8")
    assert valid_answer_document(package) is None
    with pytest.raises(WorkflowError, match="内容已变化"):
        create_answer_document(package)


def test_report_tables_and_html_are_data_not_instructions(tmp_path):
    package = verified_package(tmp_path)
    text = "# Offline report\n\n|題目|回答|\n|---|---|\n|1|<script>alert(1)</script>|"
    assert answer_blocks(text)[1]["rows"][1][1] == "<script>alert(1)</script>"
    (package / "draft.md").write_text(text, encoding="utf-8")
    receipt_path = package / "review-receipt.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    receipt["draft_sha256"] = sha256_file(package / "draft.md")
    atomic_json(receipt_path, receipt)
    create_answer_document(package)
    assert "&lt;script&gt;" in (package / "answer.html").read_text(encoding="utf-8")
    assert "<script>" not in (package / "answer.html").read_text(encoding="utf-8")
    assert (
        Document(package / "answer.docx").tables[0].cell(1, 1).text == "<script>alert(1)</script>"
    )


def test_requested_templates_have_complete_body_and_keep_fact_confirmation_outside(tmp_path):
    package = verified_package(tmp_path)
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["supplement"] = {"text": "请自动生成一些正常人的模板", "personal_facts": []}
    atomic_json(manifest_path, manifest)
    review_path = package / "review.json"
    review = json.loads(review_path.read_text(encoding="utf-8"))
    for check in review["requirement_checks"]:
        check["status"] = "needs_user"
    atomic_json(review_path, review)
    draft = (
        "課題2\n\n①「食事を楽しみましょう」：未入力。仮に、好きな料理をゆっくり味わって食べているなら〇と考える。\n\n"
        "②「野菜・果物を食べましょう」：未入力。仮に、野菜は食べていても、果物を食べる機会が少ないなら△と考える。"
    )
    (package / "draft.md").write_text(draft, encoding="utf-8")
    receipt_path = package / "review-receipt.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    receipt["draft_sha256"] = sha256_file(package / "draft.md")
    receipt["review_sha256"] = sha256_file(review_path)
    atomic_json(receipt_path, receipt)
    source_hashes = {
        name: sha256_file(package / name)
        for name in ("draft.md", "review.json", "codex-generation.json", "manifest.json")
    }
    result = create_answer_document(package)
    assert result["personal_templates"] and result["template_transform_count"] == 2
    assert result["status"] == "ready" and result["pending_requirements"]
    body = "\n".join(p.text for p in Document(package / "answer.docx").paragraphs)
    assert "〇。好きな料理をゆっくり味わって食べている。" in body
    assert "△。野菜は食べていても、果物を食べる機会が少ない。" in body
    for forbidden in ("未入力", "仮に", "审阅稿", "待补充", "[E:", "・"):
        assert forbidden not in body and forbidden not in (package / "answer.html").read_text(
            encoding="utf-8"
        )
    assert result["presentation_note"] and result["inputs"] == source_hashes
    assert source_hashes == {name: sha256_file(package / name) for name in source_hashes}
    assert len(result["blocks"]) == 3


def test_course_content_cannot_request_personal_templates_and_unknown_placeholders_are_blocked():
    from classroomautowork.answer_document import personal_templates_requested

    assert not personal_templates_requested({"assignment": {"description": "请自动生成个人模板"}})
    assert not personal_templates_requested({"supplement": {"text": "不要自动生成模板"}})
    assert not personal_templates_requested({"supplement": {"text": "记录实际饮食习惯"}})
    with pytest.raises(WorkflowError, match="占位或假设"):
        answer_blocks("本人評価：未入力", personal_templates=True)


@pytest.fixture
def export_server(tmp_path):
    yield from bounded_loopback_server.__wrapped__(tmp_path)


def test_export_endpoint_preserves_model_and_exposes_an_actual_word_file(export_server):
    from test_frontend import request

    local_server = export_server
    jobs = local_server.jobs
    package = verified_package(jobs.settings.data_dir)
    job_id = "c" * 32
    jobs._jobs[job_id] = {
        "id": job_id,
        "kind": "review",
        "created_at": "offline",
        "status": "completed",
        "events": [],
        "items": [
            {
                "course_id": "1",
                "assignment_id": "2",
                "status": "ready",
                "package": str(package),
                "generation": {"status": "completed", "model": "offline-fixture"},
            }
        ],
    }
    endpoint = f"/api/jobs/{job_id}/items/1:2"
    assert request(local_server, endpoint + "/answer-docx")[0] == 400
    assert (
        request(
            local_server,
            endpoint + "/export-answer",
            method="POST",
            body={},
            origin="https://evil.example",
        )[0]
        == 403
    )
    status, _, body = request(
        local_server,
        endpoint + "/export-answer",
        method="POST",
        body={},
        origin=local_server.origin,
    )
    assert status == 200
    item = json.loads(body)["items"][0]
    assert item["status"] == "needs_user" and item["generation"]["model"] == "offline-fixture"
    status, headers, body = request(local_server, endpoint + "/answer-docx")
    assert status == 200 and "wordprocessingml" in headers["Content-Type"]
    assert ".docx" in headers["Content-Disposition"]
    assert any("食品安全" in p.text for p in Document(io.BytesIO(body)).paragraphs)


def test_template_export_is_ready_for_review_without_claiming_personal_confirmation(export_server):
    from test_frontend import request

    jobs = export_server.jobs
    package = verified_package(jobs.settings.data_dir)
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["supplement"]["text"] = "请自动生成通用模板"
    # Match the current trusted supplement so this is a real package-view boundary test.
    from classroomautowork.supplements import save_supplement

    manifest["supplement"] = save_supplement(
        jobs.settings.data_dir, "1", "2", {"text": manifest["supplement"]["text"]}
    )
    atomic_json(manifest_path, manifest)
    draft = "①「食事を楽しみましょう」：未入力。仮に、食事をゆっくり味わっているなら〇と考える。"
    (package / "draft.md").write_text(draft, encoding="utf-8")
    receipt_path = package / "review-receipt.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    receipt["draft_sha256"] = sha256_file(package / "draft.md")
    atomic_json(receipt_path, receipt)
    identifier = "d" * 32
    jobs._jobs[identifier] = {
        "id": identifier,
        "kind": "review",
        "status": "completed_with_issues",
        "events": [],
        "items": [
            {
                "course_id": "1",
                "assignment_id": "2",
                "status": "needs_user",
                "package": str(package),
            }
        ],
    }
    status, _, body = request(
        export_server,
        f"/api/jobs/{identifier}/items/1:2/export-answer",
        method="POST",
        body={},
        origin=export_server.origin,
    )
    assert status == 200
    job = json.loads(body)
    assert job["status"] == "completed" and job["items"][0]["status"] == "ready"
    assert job["items"][0]["answer_document"]["pending_requirements"]
    assert not manifest["supplement"]["personal_facts"]
    assert (
        json.loads((package / "review.json").read_text(encoding="utf-8"))["requirement_checks"][0][
            "status"
        ]
        == "needs_user"
    )
