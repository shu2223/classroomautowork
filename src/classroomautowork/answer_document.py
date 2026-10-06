"""Readable local report files from verified answers, with no model or remote writes."""

import html
import json
import os
import re
import uuid
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .errors import WorkflowError
from .local import atomic_json, require_private_path, sha256_file, utc_now

VERSION = "answer-document-v1.2"


def answer_blocks(draft):
    """Remove explicit internal audit notes, preserving answers and unfinished fields."""
    lines = []
    for line in draft.splitlines():
        if re.match(r"^#{1,6}\s+本[机機].*(证据|證據|审核|審核|監査).*不[进入進入]", line):
            break
        lines.append(re.sub(r"\[E:[a-f0-9]{24}\]", "", line).rstrip())
    text = "\n".join(lines)
    paragraphs = re.split(r"\n\s*\n", text)
    blocks = []
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if (
            not paragraph
            or paragraph == "---"
            or paragraph.startswith(
                ("【本机审核说明】", "【本機審核說明】", "本机分析依据：", "本機分析依據：")
            )
        ):
            continue
        if paragraph.startswith("|"):
            rows = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in paragraph.splitlines()
                if not re.fullmatch(r"[| :\-]+", row)
            ]
            if rows and all(len(row) == len(rows[0]) for row in rows):
                blocks.append({"kind": "table", "rows": rows})
                continue
        heading = re.fullmatch(r"(#{1,6})\s+(.+)", paragraph)
        blocks.append(
            {
                "kind": "heading" if heading else "paragraph",
                "level": len(heading[1]) if heading else 0,
                "text": heading[2] if heading else paragraph,
            }
        )
    if not blocks:
        raise WorkflowError("没有可导出的实际答案。")
    return blocks


def _inputs(package):
    receipt = json.loads((package / "review-receipt.json").read_text(encoding="utf-8"))
    generation = json.loads((package / "codex-generation.json").read_text(encoding="utf-8"))
    if (
        generation.get("status") != "completed"
        or not receipt.get("draft_sha256")
        or sha256_file(package / "draft.md") != receipt["draft_sha256"]
        or sha256_file(package / "review.json") != receipt.get("review_sha256")
    ):
        raise WorkflowError("答案尚未完成核验或内容已变化，不能生成审阅文件。")
    return {
        name: sha256_file(package / name)
        for name in ("draft.md", "review.json", "codex-generation.json", "manifest.json")
    }


def valid_answer_document(package):
    package = require_private_path(Path(package))
    try:
        result = json.loads((package / "answer-document.json").read_text(encoding="utf-8"))
        if result.get("version") != VERSION or result["inputs"] != _inputs(package):
            return None
        for name in ("answer.docx", "answer.html"):
            path = (package / name).resolve()
            if path.parent != package or sha256_file(path) != result["files"][name]:
                return None
        return result
    except (OSError, ValueError, KeyError, TypeError, WorkflowError):
        return None


def create_answer_document(package):
    """Idempotent export; preserve source/model provenance and mark unfinished work."""
    package = require_private_path(Path(package))
    cached = valid_answer_document(package)
    if cached:
        return cached
    inputs = _inputs(package)
    manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
    review = json.loads((package / "review.json").read_text(encoding="utf-8"))
    pending = [x["requirement"] for x in review["requirement_checks"] if x["status"] != "met"]
    blocks = answer_blocks((package / "draft.md").read_text(encoding="utf-8"))
    title = manifest["assignment"].get("title", "作业答案")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Yu Mincho"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "游明朝")
        style.font.color.rgb = RGBColor(0, 0, 0)
        props = style.element.get_or_add_pPr()
        for border in list(props.findall(qn("w:pBdr"))):
            props.remove(border)
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        props.append(snap)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    document.styles["Title"].font.size = Pt(17)
    document.styles["Heading 1"].font.size = Pt(13)
    document.styles["Heading 2"].font.size = Pt(11)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].paragraph_format.space_before = Pt(10)
        document.styles[name].paragraph_format.space_after = Pt(4)
    # This is an answer review file, not a certification of assignment completion.
    banner = (
        "审阅稿：尚有内容待补充，请先在助手中补齐。"
        if pending
        else "答案审阅稿：请核对后手动提交。"
    )
    first_is_title = blocks[0]["kind"] == "heading" and blocks[0]["level"] == 1
    displayed_title = blocks[0]["text"] if first_is_title else title
    document.add_paragraph(displayed_title, "Title")
    note = document.add_paragraph(banner)
    note.runs[0].font.color.rgb = RGBColor.from_string("916000")
    rendered = [
        f"<h1>{html.escape(displayed_title)}</h1>",
        f"<p class='notice'>{html.escape(banner)}</p>",
    ]
    for block in blocks[1:] if first_is_title else blocks:
        if block["kind"] == "table":
            rows = block["rows"]
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = False
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            width = Cm(16.6 / len(rows[0]))
            for values in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, values, strict=True):
                    cell.width = width
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        for run in p.runs:
                            run.font.size = Pt(9)
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell in table.rows[0].cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shading)
            rendered.append(
                "<table>"
                + "".join(
                    "<tr>"
                    + "".join("<td>" + html.escape(value) + "</td>" for value in row)
                    + "</tr>"
                    for row in rows
                )
                + "</table>"
            )
        else:
            text = block["text"]
            level = block["level"]
            style = "Title" if level == 1 else f"Heading {min(level - 1, 3)}" if level else "Normal"
            document.add_paragraph(text, style)
            tag = f"h{min(level, 6)}" if level else "p"
            rendered.append(f"<{tag}>" + html.escape(text).replace("\n", "<br>") + f"</{tag}>")
    destination = package / "answer.docx"
    if (
        destination.resolve().parent != package
        or (package / "answer.html").resolve().parent != package
    ):
        raise WorkflowError("答案文件不能链接到其他私人文件。")
    temporary = package / ("answer-" + uuid.uuid4().hex + ".docx")
    try:
        document.save(temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    (package / "answer.html").write_text(
        "<!doctype html><html lang='ja'><meta charset='utf-8'><title>"
        + html.escape(title)
        + "</title><style>body{max-width:850px;margin:32px auto;padding:24px;font:16px/1.8 sans-serif}"
        "table{width:100%;border-collapse:collapse}td{border:1px solid #ccc;padding:8px}"
        ".notice{background:#fff4db;padding:12px}h1{font-size:24px}h2{font-size:20px}</style>"
        + "".join(rendered)
        + "</html>",
        encoding="utf-8",
    )
    result = {
        "version": VERSION,
        "created_at": utc_now(),
        "inputs": inputs,
        "title": title,
        "status": "needs_user" if pending else "ready",
        "pending_requirements": pending,
        "blocks": blocks,
        "files": {name: sha256_file(package / name) for name in ("answer.docx", "answer.html")},
        "submission": "manual_only",
    }
    atomic_json(package / "answer-document.json", result)
    return result
